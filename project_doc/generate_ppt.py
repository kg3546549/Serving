import collections 
import collections.abc
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# 테마 색상 정의 (HTML slide-theme.css 테마와 완전 동치)
LIGHT_BG = RGBColor(238, 244, 250)      # #eef4fa (슬라이드 배경)
CARD_BG = RGBColor(255, 253, 248)       # #fffdf8 (카드/박스 배경)
BORDER_COLOR = RGBColor(223, 215, 200)  # #dfd7c8 (카드 테두리)
NAVY = RGBColor(20, 34, 56)             # #142238 (주 텍스트/네이비)
MUTED = RGBColor(110, 115, 128)         # #6e7380 (본문 텍스트/그레이)
MUTED_LIGHT = RGBColor(147, 152, 163)   # #9398a3 (푸터 텍스트)

# 액센트 컬러
ACCENT_BLUE = RGBColor(136, 184, 239)   # #88b8ef
ACCENT_PURPLE = RGBColor(180, 161, 229) # #b4a1e5
ACCENT_MINT = RGBColor(125, 217, 190)   # #7dd9be
ACCENT_YELLOW = RGBColor(246, 212, 119) # #f6d477
ACCENT_PINK = RGBColor(242, 154, 172)   # #f29aac

def set_slide_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_slide_header_and_footer(slide, index, total, label, section, title, subtitle=""):
    # Header Brand (Left)
    tb_brand = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(4.0), Inches(0.4))
    tf_brand = tb_brand.text_frame
    tf_brand.word_wrap = True
    p_brand = tf_brand.paragraphs[0]
    p_brand.text = "S!!  Serving!!"
    p_brand.font.size = Pt(12)
    p_brand.font.bold = True
    p_brand.font.color.rgb = NAVY
    p_brand.font.name = 'Malgun Gothic'
    
    # Header Section (Right)
    tb_sec = slide.shapes.add_textbox(Inches(5.5), Inches(0.2), Inches(4.0), Inches(0.4))
    tf_sec = tb_sec.text_frame
    tf_sec.word_wrap = True
    p_sec = tf_sec.paragraphs[0]
    p_sec.text = section
    p_sec.font.size = Pt(11)
    p_sec.font.bold = True
    p_sec.font.color.rgb = MUTED
    p_sec.font.name = 'Malgun Gothic'
    p_sec.alignment = PP_ALIGN.RIGHT
    
    # Main Slide Title
    tb_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.45), Inches(9.0), Inches(0.6))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title
    p_title.font.size = Pt(25)
    p_title.font.bold = True
    p_title.font.color.rgb = NAVY
    p_title.font.name = 'Malgun Gothic'
    
    # Subtitle
    if subtitle:
        tb_sub = slide.shapes.add_textbox(Inches(0.5), Inches(0.95), Inches(9.0), Inches(0.4))
        tf_sub = tb_sub.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle
        p_sub.font.size = Pt(12.5)
        p_sub.font.color.rgb = MUTED
        p_sub.font.name = 'Malgun Gothic'

    # Footer Label (Left)
    tb_foot = slide.shapes.add_textbox(Inches(0.5), Inches(5.15), Inches(4.0), Inches(0.3))
    tf_foot = tb_foot.text_frame
    tf_foot.word_wrap = True
    p_foot = tf_foot.paragraphs[0]
    p_foot.text = f"Serving!! · {label}"
    p_foot.font.size = Pt(9)
    p_foot.font.bold = True
    p_foot.font.color.rgb = MUTED_LIGHT
    p_foot.font.name = 'Malgun Gothic'
    
    # Footer Index (Right)
    tb_idx = slide.shapes.add_textbox(Inches(5.5), Inches(5.15), Inches(4.0), Inches(0.3))
    tf_idx = tb_idx.text_frame
    tf_idx.word_wrap = True
    p_idx = tf_idx.paragraphs[0]
    p_idx.text = f"{str(index).zfill(2)} / {str(total).zfill(2)}"
    p_idx.font.size = Pt(9)
    p_idx.font.bold = True
    p_idx.font.color.rgb = MUTED_LIGHT
    p_idx.font.name = 'Malgun Gothic'
    p_idx.alignment = PP_ALIGN.RIGHT

def draw_card(slide, left, top, width, height, title="", text="", html_list=None, bg_color=None, border_color=None, padding=0.15):
    # Add rounded rectangle card box shape
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color or CARD_BG
    shape.line.color.rgb = border_color or BORDER_COLOR
    shape.line.width = Pt(1)
    
    # Overlay precise textbox to style contents safely
    pad = Inches(padding)
    tb = slide.shapes.add_textbox(left + pad, top + pad, width - (pad * 2), height - (pad * 2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0)
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_bottom = Inches(0)
    
    first_para = True
    if title:
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.font.name = 'Malgun Gothic'
        p.space_after = Pt(8)
        first_para = False
        
    if text:
        p_text = tf.paragraphs[0] if first_para else tf.add_paragraph()
        p_text.text = text
        p_text.font.size = Pt(10.5)
        p_text.font.color.rgb = MUTED
        p_text.font.name = 'Malgun Gothic'
        p_text.line_spacing = 1.2
        first_para = False
        
    if html_list:
        for idx, item in enumerate(html_list):
            p_list = tf.paragraphs[0] if (first_para and idx == 0) else tf.add_paragraph()
            clean_text = item.replace("<b>", "").replace("</b>", "")
            p_list.text = f"• {clean_text}"
            p_list.font.size = Pt(10)
            p_list.font.color.rgb = MUTED
            p_list.font.name = 'Malgun Gothic'
            p_list.space_before = Pt(3)
            p_list.line_spacing = 1.15
            first_para = False

def draw_node(slide, left, top, width, height, label, border_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = border_color
    shape.line.width = Pt(2)
    
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = label
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.font.name = 'Malgun Gothic'
    p.alignment = PP_ALIGN.CENTER

def create_gdd_ppt():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    slide_layout = prs.slide_layouts[6] # Blank
    
    total_slides = 7

    # ----------------------------------------------------
    # Slide 1: Cover
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    
    # Left Text Area
    tb_left = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(4.5), Inches(3.2))
    tf_left = tb_left.text_frame
    tf_left.word_wrap = True
    
    p_kicker = tf_left.paragraphs[0]
    p_kicker.text = "GAME DESIGN DOCUMENT · 2026"
    p_kicker.font.size = Pt(12)
    p_kicker.font.bold = True
    p_kicker.font.color.rgb = ACCENT_PURPLE
    p_kicker.font.name = 'Malgun Gothic'
    p_kicker.space_after = Pt(8)
    
    p_title = tf_left.add_paragraph()
    p_title.text = "Serving!!"
    p_title.font.size = Pt(46)
    p_title.font.bold = True
    p_title.font.color.rgb = NAVY
    p_title.font.name = 'Malgun Gothic'
    p_title.space_after = Pt(12)
    
    p_copy = tf_left.add_paragraph()
    p_copy.text = "클라우드 인프라를 직접 설계하고 병목을 해결하는\n웹 기반 아키텍처 설계 시뮬레이터 / 디펜스 퍼즐"
    p_copy.font.size = Pt(12.5)
    p_copy.font.color.rgb = MUTED
    p_copy.font.name = 'Malgun Gothic'
    p_copy.space_after = Pt(20)
    
    # Chips (Pills)
    chips = ["Web / PC", "Phaser 3", "React 19", "10 Waves"]
    chip_left = Inches(0.5)
    for chip in chips:
        shape_chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, chip_left, Inches(3.8), Inches(0.9), Inches(0.3))
        shape_chip.fill.solid()
        shape_chip.fill.fore_color.rgb = CARD_BG
        shape_chip.line.color.rgb = BORDER_COLOR
        shape_chip.line.width = Pt(1)
        
        tf_chip = shape_chip.text_frame
        p_chip = tf_chip.paragraphs[0]
        p_chip.text = chip
        p_chip.font.size = Pt(8.5)
        p_chip.font.bold = True
        p_chip.font.color.rgb = ACCENT_BLUE
        p_chip.font.name = 'Malgun Gothic'
        p_chip.alignment = PP_ALIGN.CENTER
        chip_left += Inches(1.0)

    # Right visual box (Navy dashboard aesthetic card matching HTML)
    shape_visual = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.3), Inches(1.0), Inches(4.2), Inches(3.6))
    shape_visual.fill.solid()
    shape_visual.fill.fore_color.rgb = NAVY
    shape_visual.line.color.rgb = NAVY
    
    # Draw Nodes inside visual box
    draw_node(slide, Inches(5.6), Inches(1.4), Inches(1.0), Inches(0.5), "INGRESS", ACCENT_BLUE)
    draw_node(slide, Inches(6.9), Inches(1.4), Inches(1.0), Inches(0.5), "ALB", ACCENT_PURPLE)
    draw_node(slide, Inches(6.25), Inches(2.4), Inches(1.0), Inches(0.5), "EC2", ACCENT_MINT)
    draw_node(slide, Inches(7.55), Inches(3.3), Inches(1.0), Inches(0.5), "RDS", ACCENT_YELLOW)
    draw_node(slide, Inches(8.2), Inches(1.4), Inches(1.0), Inches(0.5), "EGRESS", ACCENT_PINK)

    # Lines / Connectors
    # Ingress -> ALB
    line1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.6), Inches(1.63), Inches(0.3), Inches(0.03))
    line1.fill.solid(); line1.fill.fore_color.rgb = ACCENT_BLUE; line1.line.fill.background()
    # ALB -> Egress
    line2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.9), Inches(1.63), Inches(0.3), Inches(0.03))
    line2.fill.solid(); line2.fill.fore_color.rgb = ACCENT_PURPLE; line2.line.fill.background()

    # ----------------------------------------------------
    # Slide 2: 5. 게임 소개
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    add_slide_header_and_footer(slide, 2, total_slides, "기획서", "게임 소개", "5. 게임 소개 (Game Overview)", "풍선 타워 디펜스, 팩토리오, TFT, 미니 메트로가 결합된 서버 인프라 아키텍처 퍼즐")
    
    draw_card(
        slide, Inches(0.5), Inches(1.45), Inches(4.3), Inches(3.4),
        title="게임 기본 사양",
        html_list=[
            "🎮게임 제목 : Serving!! (서빙!!)",
            "🏷️게임 장르 : 서버 인프라 설계 시뮬레이터 / 디펜스 퍼즐",
            "💻플랫폼 : PC Web (브라우저 최적화)",
            "👤타깃 유저 : 네트워크/소프트웨어 학습자 및 백엔드 지망 개발자"
        ]
    )
    draw_card(
        slide, Inches(5.2), Inches(1.45), Inches(4.3), Inches(3.4),
        title="핵심 재미 및 레퍼런스",
        html_list=[
            "⭐핵심 재미 : 실시간 서버/DB 병목을 튜닝하여 극복하는 SRE 지적 쾌감",
            "📌참고 게임 : Bloons TD (배치/디펜스) + Factorio (최적 파이프라인) + TFT (상점 합성/증강) + Mini Metro (맨해튼 노선 가설)"
        ]
    )

    # ----------------------------------------------------
    # Slide 3: 6. 코어 루프
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    add_slide_header_and_footer(slide, 3, total_slides, "기획서", "코어 루프", "6. 코어 루프 (Core Loop)", "배치, 운영, 보상, 성장이 반복되며 아키텍처를 유기적으로 리팩토링하는 핵심 구조")
    
    # 4 timeline cards
    loops = [
        ("1. 플레이어 행동", "상점에서 Credits로 서버 및 모듈 장비를 구매하여 보드에 배치하고 배선함"),
        ("2. 보상 획득", "서비스 개시 후 유입되는 패킷을 정상 처리하면 Credits 재화와 XP 보상 회수"),
        ("3. 인프라 성장", " Credits로 리롤/레벨업하여 격자 보드 확장 및 동일 기기 3개 자동 합성"),
        ("4. 반복 플레이", "웨이브가 지날수록 심화되는 병목 패턴(GET, POST, SLOW, Burst)에 맞춰 리팩터링")
    ]
    card_width = Inches(2.05)
    card_height = Inches(1.6)
    left_start = Inches(0.5)
    for idx, (step_title, step_desc) in enumerate(loops):
        draw_card(slide, left_start, Inches(1.45), card_width, card_height, title=step_title, text=step_desc, padding=0.1)
        if idx < 3:
            # Draw simple arrow indicator textbox
            tb_arrow = slide.shapes.add_textbox(left_start + card_width, Inches(2.0), Inches(0.2), Inches(0.5))
            tf_arrow = tb_arrow.text_frame
            p_arrow = tf_arrow.paragraphs[0]
            p_arrow.text = "→"
            p_arrow.font.size = Pt(16)
            p_arrow.font.bold = True
            p_arrow.font.color.rgb = ACCENT_BLUE
            p_arrow.font.name = 'Malgun Gothic'
            p_arrow.alignment = PP_ALIGN.CENTER
        left_start += Inches(2.25)
        
    # 3 bottom cards (notes)
    notes = [
        ("ACTION", "Architecture Build", None),
        ("REWARD", "Credits + XP", ACCENT_YELLOW),
        ("GROWTH", "Star Upgrade & Augments", ACCENT_MINT)
    ]
    note_width = Inches(2.8)
    note_height = Inches(1.2)
    left_start_note = Inches(0.5)
    for title, text, accent in notes:
        draw_card(slide, left_start_note, Inches(3.35), note_width, note_height, title=title, text=text, border_color=accent, padding=0.1)
        left_start_note += Inches(3.1)

    # ----------------------------------------------------
    # Slide 4: 7. 세계관 소개
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    add_slide_header_and_footer(slide, 4, total_slides, "기획서", "세계관 소개", "7. 세계관 소개 (Worldview)", "신규 서비스 오픈 후 트래픽 폭주 상황 속 서버 인프라 구조를 수호해야 하는 관제 환경")
    
    worldviews = [
        ("게임의 배경", "IT 서비스 런칭 이후 기하급수적으로 폭증하는 트래픽 관제를 전개하는 백엔드 가상현실 대시보드 화면입니다.", ACCENT_BLUE),
        ("게임의 분위기", "다크 블루/네이비 톤 테마의 테크니컬 디자인으로, 실시간 광점 패킷 이동과 붉은 경고가 긴장감을 유도합니다.", None),
        ("플레이어의 목표", "실패율(Timeout, Queue Full)을 제어하고 Service HP를 사수하여 Wave 10 최종 부하 테스트를 통과하는 것입니다.", ACCENT_PINK),
        ("이 상황이 벌어진 이유", "신규 서비스의 갑작스러운 마케팅 대성공으로 사용자 동시 요청이 몰려 서버가 다운되기 일보 직전의 위기 상황입니다.", ACCENT_YELLOW),
        ("게임 속 공간의 특징", "플레이어 레벨 성장에 따라 7×4 그리드에서 시작하여 최대 13×6 그리드까지 물리적 가설 영역이 점진적으로 확장됩니다.", ACCENT_MINT),
        ("플레이어의 역할", "제한된 연결선 길이(링크 예산)와 격자 면적 안에서 최대 가용성을 확보하는 현실적 클라우드 아키텍트의 위치입니다.", ACCENT_PURPLE)
    ]
    card_width = Inches(2.8)
    card_height = Inches(1.6)
    for idx, (title, desc, accent) in enumerate(worldviews):
        row = idx // 3
        col = idx % 3
        left = Inches(0.5) + col * Inches(3.1)
        top = Inches(1.45) + row * Inches(1.8)
        draw_card(slide, left, top, card_width, card_height, title=title, text=desc, border_color=accent, padding=0.1)

    # ----------------------------------------------------
    # Slide 5: 8. 캐릭터 소개
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    add_slide_header_and_footer(slide, 5, total_slides, "기획서", "캐릭터 소개", "8. 캐릭터 소개 (Characters)", "인프라 설계를 담당하는 아키텍트와 병목을 일으키는 4종의 요청 유형(적) 및 최종 단계")
    
    # Native Table shape
    rows, cols = 5, 4
    left, top, width, height = Inches(0.5), Inches(1.45), Inches(9.0), Inches(3.3)
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    
    # Column widths
    table.columns[0].width = Inches(1.1)
    table.columns[1].width = Inches(1.7)
    table.columns[2].width = Inches(3.4)
    table.columns[3].width = Inches(2.8)
    
    table_data = [
        ["구분", "대상", "역할 및 위협 패턴", "대응 전략 및 특징"],
        ["주인공", "인프라 아키텍트", "상점 경제 관리, 서버 배치 및 데이터/트래픽 링크 가설", "예산과 거리 제약을 극복하며 병목 구간 튜닝"],
        ["적 캐릭터", "GET / POST 요청", "DB 읽기 부하 집중 패킷 / DB 쓰기 집중 부하를 유도하여 쓰기 병목 유발 패킷", "RDS Replica 배치를 통한 읽기 분산 / DB 업그레이드"],
        ["적 캐릭터", "SLOW / Burst 요청", "인덱싱 없는 장기 쿼리 (DB 큐 마비) / 짧은 주기로 동시 유입되어 서버 큐 폭발 유발", "ElastiCache Redis 도입 / SQS 큐 버퍼 탑재 완충"],
        ["보스", "부하 테스트 (Wave 10)", "GET, POST, SLOW, Burst가 복합적으로 일시에 쏟아지는 임계치 테스트", "합성(2성/3성) 및 무작위 증강 시너지 최종 검증"]
    ]
    for row_idx, row in enumerate(table_data):
        for col_idx, text in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = text
            # Format cell fill
            cell.fill.solid()
            if row_idx == 0:
                cell.fill.fore_color.rgb = RGBColor(241, 246, 251) # --navy light background
            else:
                cell.fill.fore_color.rgb = CARD_BG
            
            # Text formatting
            p = cell.text_frame.paragraphs[0]
            p.font.name = 'Malgun Gothic'
            p.font.size = Pt(9.5)
            if row_idx == 0:
                p.font.bold = True
                p.font.color.rgb = NAVY
            else:
                p.font.color.rgb = MUTED

    # ----------------------------------------------------
    # Slide 6: 9. 컨셉 아트
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    add_slide_header_and_footer(slide, 6, total_slides, "기획서", "컨셉 아트", "9. 컨셉 아트 (Concept Art)", "게임 비주얼 방향성 및 스크린샷 이미지 가설 영역")
    
    # Left placeholder card
    draw_card(
        slide, Inches(0.5), Inches(1.45), Inches(4.3), Inches(3.4),
        title="[컨셉 아트 이미지 영역]",
        text="본 칸은 사용자가 직접 실제 개발 스크린샷 및 UI 아이콘 이미지를 파워포인트 내에서 바로 삽입하는 이미지 플레이스홀더 칸입니다.",
        bg_color=RGBColor(240, 237, 230), border_color=RGBColor(200, 190, 180)
    )
    # Right explanation card
    draw_card(
        slide, Inches(5.2), Inches(1.45), Inches(4.3), Inches(3.4),
        title="Visual Direction",
        html_list=[
            "다크 블루 테마 (#142238) 기반의 정갈한 네오 브루탈리즘 테크 HUD 스타일",
            "실시간 패킷 광점 이동 경로(파랑: 요청, 보라: 응답, 노랑: DB) 시각화",
            "과부하 큐에 직관적인 빨간색 위험 경고 노출"
        ]
    )
    # Draw Swatches inside right card
    swatches = [NAVY, ACCENT_BLUE, ACCENT_PURPLE, ACCENT_YELLOW, ACCENT_MINT, ACCENT_PINK]
    start_x = Inches(5.5)
    for color in swatches:
        shape_circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, start_x, Inches(4.0), Inches(0.4), Inches(0.4))
        shape_circle.fill.solid()
        shape_circle.fill.fore_color.rgb = color
        shape_circle.line.color.rgb = color
        start_x += Inches(0.55)

    # ----------------------------------------------------
    # Slide 7: Summary Cover
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    
    # Header Brand (Left)
    tb_brand = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(4.0), Inches(0.4))
    p_brand = tb_brand.text_frame.paragraphs[0]
    p_brand.text = "S!!  Serving!!"
    p_brand.font.size = Pt(12)
    p_brand.font.bold = True
    p_brand.font.color.rgb = NAVY
    p_brand.font.name = 'Malgun Gothic'

    # Content Box
    tb_sum = slide.shapes.add_textbox(Inches(1.0), Inches(1.4), Inches(8.0), Inches(3.0))
    tf_sum = tb_sum.text_frame
    tf_sum.word_wrap = True
    
    p_kicker = tf_sum.paragraphs[0]
    p_kicker.text = "DESIGN SUMMARY"
    p_kicker.font.size = Pt(13)
    p_kicker.font.bold = True
    p_kicker.font.color.rgb = ACCENT_PURPLE
    p_kicker.font.name = 'Malgun Gothic'
    p_kicker.alignment = PP_ALIGN.CENTER
    p_kicker.space_after = Pt(8)
    
    p_title = tf_sum.add_paragraph()
    p_title.text = "Serving!!"
    p_title.font.size = Pt(42)
    p_title.font.bold = True
    p_title.font.color.rgb = NAVY
    p_title.font.name = 'Malgun Gothic'
    p_title.alignment = PP_ALIGN.CENTER
    p_title.space_after = Pt(12)
    
    p_copy = tf_sum.add_paragraph()
    p_copy.text = "보이지 않던 서버 인프라 설계 지식을\n직관적인 배치, 배선, 병목 대응 플레이 행동으로 이해하는 게임"
    p_copy.font.size = Pt(13.5)
    p_copy.font.color.rgb = MUTED
    p_copy.font.name = 'Malgun Gothic'
    p_copy.alignment = PP_ALIGN.CENTER
    p_copy.space_after = Pt(20)

    # Save
    SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
    output_path = os.path.join(SCRIPT_DIR, "기획서.pptx")
    prs.save(output_path)
    print(f"Saved: {output_path}")

def create_dev_ppt():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    slide_layout = prs.slide_layouts[6]
    
    total_slides = 5

    # ----------------------------------------------------
    # Slide 1: Cover
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    
    tb_left = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(8.0), Inches(3.0))
    tf_left = tb_left.text_frame
    tf_left.word_wrap = True
    
    p_kicker = tf_left.paragraphs[0]
    p_kicker.text = "DEVELOPMENT PLAN · 2026"
    p_kicker.font.size = Pt(12)
    p_kicker.font.bold = True
    p_kicker.font.color.rgb = ACCENT_PURPLE
    p_kicker.font.name = 'Malgun Gothic'
    p_kicker.alignment = PP_ALIGN.CENTER
    p_kicker.space_after = Pt(8)
    
    p_title = tf_left.add_paragraph()
    p_title.text = "Serving!! 개발 계획서"
    p_title.font.size = Pt(42)
    p_title.font.bold = True
    p_title.font.color.rgb = NAVY
    p_title.font.name = 'Malgun Gothic'
    p_title.alignment = PP_ALIGN.CENTER
    p_title.space_after = Pt(12)
    
    p_copy = tf_left.add_paragraph()
    p_copy.text = "화면 구성부터 핵심 기능 구현, 우선순위 및 개발 진행 기준"
    p_copy.font.size = Pt(13)
    p_copy.font.color.rgb = MUTED
    p_copy.font.name = 'Malgun Gothic'
    p_copy.alignment = PP_ALIGN.CENTER

    # ----------------------------------------------------
    # Slide 2: 1. 화면별 개발 순서
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    add_slide_header_and_footer(slide, 2, total_slides, "개발 계획서", "화면별 개발 순서", "1. 화면별 개발 순서", "각 화면을 독립적으로 완성한 뒤 상태 공유와 전환 흐름을 연결한다.")
    
    # 6 timeline cards
    screens = [
        ("1. 메인 메뉴", "소개 · 시작 · 조작 가이드"),
        ("2. 게임 보드", "격자 · 카메라 · 노드 배치"),
        ("3. 상점·인벤", "구매 · 리롤 · 드래그 스냅"),
        ("4. 링크 편집", "포트 제약 · Manhattan 거리"),
        ("5. 웨이브 HUD", "HP · 큐 게이지 · 패킷 이동"),
        ("6. 결과·증강", "보상 credits · 합성 · 증강")
    ]
    card_width = Inches(1.35)
    card_height = Inches(1.6)
    left_start = Inches(0.5)
    for idx, (step_title, step_desc) in enumerate(screens):
        draw_card(slide, left_start, Inches(1.45), card_width, card_height, title=step_title, text=step_desc, padding=0.08)
        if idx < 5:
            tb_arrow = slide.shapes.add_textbox(left_start + card_width, Inches(2.0), Inches(0.15), Inches(0.5))
            tf_arrow = tb_arrow.text_frame
            p_arrow = tf_arrow.paragraphs[0]
            p_arrow.text = "→"
            p_arrow.font.size = Pt(13)
            p_arrow.font.bold = True
            p_arrow.font.color.rgb = ACCENT_BLUE
            p_arrow.font.name = 'Malgun Gothic'
            p_arrow.alignment = PP_ALIGN.CENTER
        left_start += Inches(1.5)
        
    # 3 bottom cards (notes)
    notes = [
        ("화면 단위 기준", "레이아웃 → 입력 → 상태 반영 → 전환", None),
        ("통합 기준", "React와 Phaser가 동일 Zustand 상태 사용", ACCENT_PURPLE),
        ("검증 기준", "전체 플레이 흐름을 브라우저에서 연속 재현", ACCENT_MINT)
    ]
    note_width = Inches(2.8)
    note_height = Inches(1.2)
    left_start_note = Inches(0.5)
    for title, text, accent in notes:
        draw_card(slide, left_start_note, Inches(3.35), note_width, note_height, title=title, text=text, border_color=accent, padding=0.1)
        left_start_note += Inches(3.1)

    # ----------------------------------------------------
    # Slide 3: 2. 기능별 개발 순서
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    add_slide_header_and_footer(slide, 3, total_slides, "개발 계획서", "기능별 개발 순서", "2. 기능별 개발 순서", "핵심 조작 가설에서부터 시뮬레이션 루프 및 합성/증강 통합까지 구현 순서")
    
    draw_card(
        slide, Inches(0.5), Inches(1.45), Inches(4.3), Inches(3.4),
        title="프론트엔드 조작 및 HUD 가설",
        html_list=[
            "그리드 보드 구축 : Phaser 타일 렌더러와 마우스 줌/스크롤 카메라 완성",
            "React UI 렌더링 : Zustand 연동 경제/HP 정보 상태판 및 카드 상점, 인벤토리 구성",
            "노드 드래그 배치 : 상점 기기를 드래그하여 보드판 셀 중심에 스냅 드롭 처리"
        ]
    )
    draw_card(
        slide, Inches(5.2), Inches(1.45), Inches(4.3), Inches(3.4),
        title="시뮬레이션 로직 및 합성/증강 통합",
        html_list=[
            "링크 연결선 시스템 : Shift+드래그 연결, 포트 규칙 검사 및 Manhattan 한계 연장 검증",
            "실시간 시뮬레이터 : Tick 기반 패킷 인스턴스 생성, Tween 물리 이동 및 컴포넌트 큐 게이지 연출",
            "오토배틀러 경제 통합 : 동일 기기 3성 자동 합성 판정 및 무작위 3종 증강 능력치 스펙 연동"
        ],
        border_color=ACCENT_PURPLE
    )

    # ----------------------------------------------------
    # Slide 4: 3. 우선순위 정리
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    add_slide_header_and_footer(slide, 4, total_slides, "개발 계획서", "우선순위 정리", "3. 우선순위 정리", "먼저 구현해야 할 핵심 기능과 나중에 추가할 기능을 명확히 구분")
    
    draw_card(
        slide, Inches(0.5), Inches(1.45), Inches(2.8), Inches(3.4),
        title="우선순위 1순위 (핵심 기능)",
        html_list=[
            "Phaser 보드 및 카메라 스크롤 조작",
            "React HUD 화면 및 노드 배치 스냅",
            "맨해튼 거리 및 포트 규칙 링크 검증"
        ]
    )
    draw_card(
        slide, Inches(3.6), Inches(1.45), Inches(2.8), Inches(3.4),
        title="우선순위 2순위 (게임 규칙)",
        html_list=[
            "실시간 Tick 시뮬레이터 및 패킷 이동",
            "상점 리롤 재화 소비 및 상태 관리",
            "동일 장비 3개 수집 시 자동 합성"
        ],
        border_color=ACCENT_PURPLE
    )
    draw_card(
        slide, Inches(6.7), Inches(1.45), Inches(2.8), Inches(3.4),
        title="우선순위 3순위 (연출/폴리싱)",
        html_list=[
            "합성 시 3종 무작위 증강 선택",
            "Glow/Particle 발광 광점 및 궤적 VFX",
            "상태별 사운드 리소스 효과음 통합"
        ],
        border_color=ACCENT_MINT
    )

    # ----------------------------------------------------
    # Slide 5: 4. 개발 진행 기준
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    add_slide_header_and_footer(slide, 5, total_slides, "개발 계획서", "개발 진행 기준", "4. 개발 진행 기준 (DoD)", "각 단계가 완료되었는지 확인하고 코드베이스를 검증하는 객관적인 기준 명시")
    
    draw_card(
        slide, Inches(0.5), Inches(1.45), Inches(2.8), Inches(3.4),
        title="기능 동작 검증",
        text="Vite 로컬 개발 환경(localhost:5173) 상에서 화면 레이아웃 깨짐이 없고 드래그 배치 및 링크 가설이 부드럽게 재현되어야 합니다.",
        border_color=ACCENT_BLUE
    )
    draw_card(
        slide, Inches(3.6), Inches(1.45), Inches(2.8), Inches(3.4),
        title="자동화 테스트 통과",
        text="거리 계산, 포트 제약 조건 검증, 상점 가챠 확률, 합성 판정 등 핵심 연산 로직에 대한 Vitest 단위 테스트(npm test)가 100% 통과(GREEN)해야 합니다.",
        border_color=ACCENT_PURPLE
    )
    draw_card(
        slide, Inches(6.7), Inches(1.45), Inches(2.8), Inches(3.4),
        title="통합 시나리오 검증",
        text="장비 조달 ➡️ 링크 가설 ➡️ 시뮬레이션 ➡️ 보상/합성 ➡️ HP 유지 사수 ➡️ Wave 10 클리어까지의 게임 루프가 에러 없이 연속 수행되어야 합니다.",
        border_color=ACCENT_MINT
    )

    # 상대 경로로 저장되도록 수정
    SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
    output_path = os.path.join(SCRIPT_DIR, "개발계획서.pptx")
    prs.save(output_path)
    print(f"Saved: {output_path}")

if __name__ == '__main__':
    create_gdd_ppt()
    create_dev_ppt()
    print("PPT files generated successfully!")
